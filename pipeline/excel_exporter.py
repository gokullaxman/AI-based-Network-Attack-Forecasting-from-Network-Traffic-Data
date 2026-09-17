"""
Excel Export Engine for Network Attack Forecasting System
Uses openpyxl to generate professionally styled, formatted .xlsx workbooks
logging sequential forecast history with SHAP attributions and defense recommendations.
"""

import io
from datetime import datetime
from typing import List, Dict, Any

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


HEADERS = [
    "Timestamp / Window Index",
    "Source",
    "Current Stage",
    "Predicted Next Stage",
    "Confidence (%)",
    "Dynamic Risk Score (0-100)",
    "Risk Level",
    "Top SHAP Feature 1",
    "Top SHAP Feature 2",
    "Top SHAP Feature 3",
    "Recommended Defenses"
]

# Risk Level Color Fills (Light visual scanning)
RISK_FILLS = {
    "HIGH": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
    "MEDIUM": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
    "LOW": PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
}

RISK_FONTS = {
    "HIGH": Font(name="Arial", size=10, bold=True, color="B91C1C"),
    "MEDIUM": Font(name="Arial", size=10, bold=True, color="B45309"),
    "LOW": Font(name="Arial", size=10, bold=True, color="047857")
}


def build_forecast_workbook(records: List[Dict[str, Any]]) -> openpyxl.Workbook:
    """
    Builds an openpyxl Workbook with one row per forecast event.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attack Forecast Log"
    
    # Freeze header row
    ws.freeze_panes = "A2"
    ws.views.sheetView[0].showGridLines = True

    # Styling definitions
    header_font = Font(name="Arial", size=11, bold=True, color="1E293B")
    header_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    body_font = Font(name="Arial", size=10, color="0F172A")
    body_align_left = Alignment(horizontal="left", vertical="center")
    body_align_center = Alignment(horizontal="center", vertical="center")
    body_align_wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)
    
    thin_border_side = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side
    )

    # 1. Write Header Row
    ws.append(HEADERS)
    ws.row_dimensions[1].height = 28
    
    for col_num in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = cell_border

    # 2. Write Data Rows
    current_time_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    for row_idx, rec in enumerate(records, start=2):
        ws.row_dimensions[row_idx].height = 50  # Room for multi-line defense recommendations
        
        # Format Timestamp / Window Index
        ts = rec.get("timestamp") or f"{current_time_str} [Win #{row_idx-1}]"
        
        # Source
        source = rec.get("data_source", "Synthetic Flow Stream")
        
        # Stages & Metrics
        current_stage = rec.get("current_stage", "Unknown")
        predicted_next = rec.get("predicted_next_stage", "Unknown")
        confidence = rec.get("confidence", 0)
        risk_score = rec.get("risk_score", 0)
        risk_level = rec.get("risk_level", "LOW")

        # Top 3 SHAP Features
        shap_ranking = rec.get("shap", {}).get("feature_ranking", [])
        top_shap_1 = ""
        top_shap_2 = ""
        top_shap_3 = ""
        if len(shap_ranking) > 0:
            top_shap_1 = f"{shap_ranking[0]['feature']} ({shap_ranking[0].get('attribution', 0):+.4f})"
        if len(shap_ranking) > 1:
            top_shap_2 = f"{shap_ranking[1]['feature']} ({shap_ranking[1].get('attribution', 0):+.4f})"
        if len(shap_ranking) > 2:
            top_shap_3 = f"{shap_ranking[2]['feature']} ({shap_ranking[2].get('attribution', 0):+.4f})"

        # Recommended Defenses (Newline-separated in one cell)
        defenses_text = ""
        rec_data = rec.get("recommendation", {})
        options = rec_data.get("options", [])
        if options:
            lines = []
            for opt in options:
                p = opt.get("priority", 1)
                act = opt.get("action", "")
                rat = opt.get("rationale", "")
                lines.append(f"[{p}] {act}\n    Rationale: {rat}")
            defenses_text = "\n".join(lines)
        elif rec_data.get("title"):
            defenses_text = f"{rec_data.get('title')}: {rec_data.get('description', '')}"

        row_values = [
            ts,
            source,
            current_stage,
            predicted_next,
            confidence,
            risk_score,
            risk_level,
            top_shap_1,
            top_shap_2,
            top_shap_3,
            defenses_text
        ]
        ws.append(row_values)

        # Apply cell-by-cell formatting
        for col_num in range(1, len(row_values) + 1):
            cell = ws.cell(row=row_idx, column=col_num)
            cell.font = body_font
            cell.border = cell_border
            
            # Alignments
            if col_num in (1, 2):
                cell.alignment = body_align_left
            elif col_num in (3, 4, 5, 6, 7):
                cell.alignment = body_align_center
            elif col_num in (8, 9, 10):
                cell.alignment = body_align_left
            elif col_num == 11:
                cell.alignment = body_align_wrap

            # Conditional Fill for Risk Level (Column 7)
            if col_num == 7 and risk_level in RISK_FILLS:
                cell.fill = RISK_FILLS[risk_level]
                cell.font = RISK_FONTS[risk_level]

    # 3. Explicit Column Widths (Prevent any truncated columns)
    col_widths = {
        1: 26,  # Timestamp
        2: 24,  # Source
        3: 18,  # Current Stage
        4: 20,  # Predicted Next
        5: 15,  # Confidence
        6: 18,  # Risk Score
        7: 15,  # Risk Level
        8: 25,  # SHAP 1
        9: 25,  # SHAP 2
        10: 25, # SHAP 3
        11: 58  # Recommended Defenses
    }
    for col_idx, width in col_widths.items():
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    return wb


def export_forecasts_to_bytes(records: List[Dict[str, Any]]) -> bytes:
    """Generates the openpyxl workbook and returns raw bytes for file download."""
    wb = build_forecast_workbook(records)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
